from __future__ import annotations

import json
import re
from io import BytesIO

from docx import Document

from sovereign_workbench.core.llm.provider import Completion
from sovereign_workbench.core.llm.registry import get_registry


def _patch_model(monkeypatch, *, plan: dict[str, object]) -> None:
    async def fake_plan(*args, **kwargs):
        spec = args[1]
        return Completion(text=json.dumps(plan), model_id=spec.id)

    async def fake_answer(*args, **kwargs):
        spec = args[1]
        return Completion(text="Local run completed with a verified result.", model_id=spec.id)

    monkeypatch.setattr(
        "sovereign_workbench.core.agent.planner.run_chat",
        fake_plan,
    )
    monkeypatch.setattr(
        "sovereign_workbench.core.agent.loop.run_chat",
        fake_answer,
    )


def _events(response_text: str) -> list[tuple[str, dict[str, object]]]:
    return [
        (event_type, json.loads(payload))
        for event_type, payload in re.findall(r"event: ([a-z_]+)\ndata: (.+?)\n\n", response_text)
    ]


def test_inspection_attachment_to_downloadable_docx(client, monkeypatch):
    _patch_model(
        monkeypatch,
        plan={
            "steps": [
                {"title": "Inspect attached report", "tool": "list_files", "kind": "tool_call"},
                {
                    "title": "Write approval note",
                    "tool": "create_approval_note",
                    "kind": "tool_call",
                    "args": {
                        "source_paths": ["inspection.txt"],
                        "subject": "Inspection report R-17",
                        "recommendation": "Approve replacement before restart.",
                    },
                },
            ]
        },
    )
    upload = client.post(
        "/api/documents",
        files={
            "file": (
                "inspection.txt",
                b"Inspection R-17\nRelief valve shows corrosion.\nReplace before restart.",
                "text/plain",
            )
        },
    )
    assert upload.status_code == 201
    document_id = upload.json()["document_id"]
    assert upload.json()["chunks"] == 1

    created = client.post(
        "/api/tasks",
        json={
            "prompt": (
                "Read the attached inspection report and draft an approval note as a Word DOCX."
            ),
            "attachments": [document_id],
        },
    )
    assert created.status_code == 201
    task_id = created.json()["id"]
    assert created.json()["attachments"] == [document_id]
    assert created.json()["model_id"]

    streamed = client.get(f"/api/tasks/{task_id}/events")
    assert streamed.status_code == 200
    events = _events(streamed.text)
    event_names = [name for name, _ in events]
    assert "plan" in event_names
    assert "tool_result" in event_names
    assert "artifact" in event_names
    assert event_names[-1] == "done"
    artifact_payload = _event_data(next(payload for name, payload in events if name == "artifact"))
    assert artifact_payload["filename"] == "approval_note.docx"

    finished = client.get(f"/api/tasks/{task_id}")
    assert finished.status_code == 200
    body = finished.json()
    assert body["status"] == "done"
    assert body["artifacts"][0]["kind"] == "docx"
    assert body["steps"]

    artifact_id = body["artifacts"][0]["id"]
    download = client.get(f"/api/artifacts/{artifact_id}")
    assert download.status_code == 200
    assert download.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert download.content[:2] == b"PK"
    content = "\n".join(
        paragraph.text for paragraph in Document(BytesIO(download.content)).paragraphs
    )
    assert "Relief valve shows corrosion" in content
    assert "Approve replacement before restart" in content
    assert "inspection.txt" in content

    audit = client.get("/api/audit", params={"type": "artifact_created"})
    assert audit.status_code == 200
    assert any(task_id in (event["detail"] + json.dumps(event["data"])) for event in audit.json())


def test_code_task_runs_in_sandbox_and_returns_source_artifact(client, monkeypatch):
    _patch_model(
        monkeypatch,
        plan={
            "steps": [
                {
                    "title": "Run verification code",
                    "tool": "run_python",
                    "kind": "tool_call",
                    "args": {"code": "print(2 + 2)"},
                }
            ]
        },
    )
    created = client.post(
        "/api/tasks",
        json={
            "prompt": "Run this Python code in the sandbox and verify it.",
            "options": {"run_code": True},
        },
    )
    assert created.status_code == 201
    task_id = created.json()["id"]

    streamed = client.get(f"/api/tasks/{task_id}/events")
    events = _events(streamed.text)
    result = _event_data(next(payload for name, payload in events if name == "tool_result"))
    assert result["ok"] is True
    assert "4" in result["output"]
    assert any(name == "done" for name, _ in events)

    body = client.get(f"/api/tasks/{task_id}").json()
    assert body["status"] == "done"
    assert body["artifacts"][0]["kind"] == "code"
    code_download = client.get(f"/api/artifacts/{body['artifacts'][0]['id']}")
    assert code_download.status_code == 200
    assert code_download.content == b"print(2 + 2)"


def test_empty_attachment_has_explicit_ocr_failure(client, monkeypatch):
    _patch_model(
        monkeypatch,
        plan={
            "steps": [
                {
                    "title": "Create approval note",
                    "tool": "create_approval_note",
                    "kind": "tool_call",
                }
            ]
        },
    )
    upload = client.post(
        "/api/documents",
        files={"file": ("blank.txt", b"", "text/plain")},
    )
    assert upload.status_code == 201
    assert upload.json()["status"] == "empty"

    created = client.post(
        "/api/tasks",
        json={
            "prompt": "Draft an approval note from this inspection report.",
            "attachments": [upload.json()["document_id"]],
        },
    )
    task_id = created.json()["id"]
    streamed = client.get(f"/api/tasks/{task_id}/events")
    events = _events(streamed.text)
    error = _event_data(next(payload for name, payload in events if name == "error"))
    assert error["code"] == "OCR_NO_TEXT"
    assert not any(name == "done" for name, _ in events)
    assert client.get(f"/api/tasks/{task_id}").json()["status"] == "error"


def test_scanned_attachment_uses_local_ocr_and_vision_route(client, monkeypatch):
    def fake_ocr(path):
        return "Scanned inspection: tank flange cracked; replace before startup."

    monkeypatch.setattr("sovereign_workbench.core.multimodal.ocr.extract_text", fake_ocr)
    _patch_model(monkeypatch, plan={"steps": []})
    upload = client.post(
        "/api/documents",
        files={"file": ("scan.png", b"mock-image", "image/png")},
    )
    assert upload.status_code == 201
    created = client.post(
        "/api/tasks",
        json={
            "prompt": "Draft an approval note from the attached report.",
            "attachments": [upload.json()["document_id"]],
        },
    )
    assert created.status_code == 201
    assert created.json()["model_id"] == "vision-min"
    task_id = created.json()["id"]
    events = _events(client.get(f"/api/tasks/{task_id}/events").text)
    assert any(name == "done" for name, _ in events)
    assert client.get(f"/api/tasks/{task_id}").json()["artifacts"][0]["kind"] == "docx"


def test_no_model_and_no_tool_reports_error(client, monkeypatch):
    async def unavailable(*args, **kwargs):
        raise RuntimeError("local model is offline")

    monkeypatch.setattr("sovereign_workbench.core.agent.planner.run_chat", unavailable)
    monkeypatch.setattr("sovereign_workbench.core.agent.loop.run_chat", unavailable)
    created = client.post("/api/tasks", json={"prompt": "Summarise the provided information."})
    assert created.status_code == 201
    task_id = created.json()["id"]
    events = _events(client.get(f"/api/tasks/{task_id}/events").text)
    errors = [_event_data(payload) for name, payload in events if name == "error"]
    assert errors[0]["code"] == "MODEL_UNAVAILABLE"
    assert not any(name == "done" for name, _ in events)
    assert client.get(f"/api/tasks/{task_id}").json()["status"] == "error"


def test_rejects_missing_attachment_and_denied_code_execution(client, monkeypatch):
    invalid = client.post(
        "/api/tasks",
        json={
            "prompt": "Read this report",
            "attachments": ["00000000-0000-0000-0000-000000000001"],
        },
    )
    assert invalid.status_code == 422
    assert "not found" in invalid.json()["error"]["message"]

    _patch_model(
        monkeypatch,
        plan={"steps": [{"title": "Run code", "tool": "run_python", "args": {"code": "print(1)"}}]},
    )
    task_id = client.post("/api/tasks", json={"prompt": "Run code to verify a sum"}).json()["id"]
    events = _events(client.get(f"/api/tasks/{task_id}/events").text)
    errors = [_event_data(payload) for name, payload in events if name == "error"]
    assert errors and "not permitted" in errors[0]["message"]
    assert client.get(f"/api/tasks/{task_id}").json()["status"] == "error"
    audit = client.get("/api/audit", params={"type": "tool_denied"}).json()
    assert any(event["data"]["tool"] == "run_python" for event in audit)


def test_deleted_attachment_and_cancelled_task_have_terminal_states(client, monkeypatch):
    upload = client.post(
        "/api/documents",
        files={"file": ("temporary.txt", b"Inspection note", "text/plain")},
    )
    document_id = upload.json()["document_id"]
    created = client.post(
        "/api/tasks",
        json={"prompt": "Draft an approval note", "attachments": [document_id]},
    )
    task_id = created.json()["id"]
    assert client.delete(f"/api/documents/{document_id}").json() == {"ok": True}
    events = _events(client.get(f"/api/tasks/{task_id}/events").text)
    error = _event_data(next(payload for name, payload in events if name == "error"))
    assert error["code"] == "ATTACHMENT_MISSING"
    assert client.get(f"/api/tasks/{task_id}").json()["status"] == "error"

    queued = client.post("/api/tasks", json={"prompt": "Summarise something"}).json()
    cancelled = client.post(f"/api/tasks/{queued['id']}/cancel")
    assert cancelled.status_code == 200
    assert client.get(f"/api/tasks/{queued['id']}").json()["status"] == "cancelled"
    replay = _events(client.get(f"/api/tasks/{queued['id']}/events").text)
    assert _event_data(replay[-1][1])["code"] == "TASK_CANCELLED"


def test_approval_note_falls_back_to_extraction_without_inference(client, monkeypatch):
    registry = get_registry()
    monkeypatch.setattr(registry, "probed", True)
    upload = client.post(
        "/api/documents",
        files={
            "file": (
                "offline-report.txt",
                b"Inspection finding: damaged safety seal; replace immediately.",
                "text/plain",
            )
        },
    )
    created = client.post(
        "/api/tasks",
        json={
            "prompt": "Draft an approval note from this inspection report.",
            "attachments": [upload.json()["document_id"]],
        },
    )
    assert created.status_code == 201
    assert created.json()["model_id"] is None
    task_id = created.json()["id"]
    streamed = client.get(f"/api/tasks/{task_id}/events")
    events = _events(streamed.text)
    assert any(name == "artifact" for name, _ in events)
    assert events[-1][0] == "done"
    assert client.get(f"/api/tasks/{task_id}").json()["status"] == "done"
    assert client.get("/api/health").json()["models_ready"] is False


def test_approval_note_rejects_incomplete_sources_and_ungrounded_findings(client, monkeypatch):
    _patch_model(
        monkeypatch,
        plan={
            "steps": [
                {
                    "title": "Create approval note",
                    "tool": "create_approval_note",
                    "args": {"findings": ["No equipment defects were found"]},
                }
            ]
        },
    )
    good = client.post(
        "/api/documents",
        files={"file": ("valve.txt", b"Valve seat cracked; replace it.", "text/plain")},
    ).json()["document_id"]
    blank = client.post(
        "/api/documents",
        files={"file": ("unreadable.txt", b"", "text/plain")},
    ).json()["document_id"]
    created = client.post(
        "/api/tasks",
        json={"prompt": "Draft an approval note", "attachments": [good, blank]},
    )
    task_id = created.json()["id"]
    events = _events(client.get(f"/api/tasks/{task_id}/events").text)
    assert _event_data(events[-1][1])["code"] == "OCR_INCOMPLETE"
    assert client.get(f"/api/tasks/{task_id}").json()["artifacts"] == []

    created = client.post(
        "/api/tasks",
        json={"prompt": "Draft an approval note", "attachments": [good]},
    )
    task_id = created.json()["id"]
    client.get(f"/api/tasks/{task_id}/events")
    artifact = client.get(f"/api/tasks/{task_id}").json()["artifacts"][0]
    body = "\n".join(
        paragraph.text
        for paragraph in Document(
            BytesIO(client.get(f"/api/artifacts/{artifact['id']}").content)
        ).paragraphs
    )
    assert "Valve seat cracked" in body
    assert "No equipment defects" not in body


def _event_data(payload: dict[str, object]) -> dict[str, object]:
    data = payload.get("data")
    return data if isinstance(data, dict) else payload
