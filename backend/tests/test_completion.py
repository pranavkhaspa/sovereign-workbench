from __future__ import annotations

import asyncio
import json
import shutil
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from sovereign_workbench.core.deliverables.calculations import calculate
from sovereign_workbench.core.deliverables.xlsx import write_xlsx
from sovereign_workbench.core.llm.provider import Completion
from sovereign_workbench.core.rag.extract import extract_pages
from sovereign_workbench.core.sandbox.subprocess_sandbox import SubprocessSandbox


def patch_plan(monkeypatch, steps):
    async def plan(*args, **kwargs):
        return Completion(text=json.dumps({"steps": steps}), model_id=args[1].id)

    async def answer(*args, **kwargs):
        yield "Finished locally."

    monkeypatch.setattr("sovereign_workbench.core.agent.planner.run_chat", plan)
    monkeypatch.setattr("sovereign_workbench.core.agent.loop.run_stream", answer)


@pytest.mark.parametrize(
    "tool,args,kind",
    [
        ("create_document", {"title": "Review", "paragraphs": ["Valve replaced."]}, "docx"),
        (
            "create_presentation",
            {"title": "Review", "slides": [{"title": "Findings", "body": "Valve replaced."}]},
            "pptx",
        ),
        ("create_spreadsheet", {"rows": [["Cost"], [12], [18], ["=SUM(A2:A3)"]]}, "xlsx"),
    ],
)
def test_office_artifacts_are_valid_downloads(client, monkeypatch, tool, args, kind):
    patch_plan(monkeypatch, [{"title": "Create deliverable", "tool": tool, "args": args}])
    response = client.post("/api/chat", json={"prompt": "Create a deliverable"})
    assert "event: done" in response.text, response.text
    artifact = client.get("/api/artifacts").json()[0]
    assert artifact["kind"] == kind
    content = client.get(f"/api/artifacts/{artifact['id']}/download").content
    if kind == "docx":
        assert Document(BytesIO(content)).paragraphs[1].text == "Valve replaced."
    elif kind == "pptx":
        deck = Presentation(BytesIO(content))
        assert len(deck.slides) == 2
        assert deck.slides[1].shapes.title.text == "Findings"
    else:
        workbook = load_workbook(BytesIO(content), data_only=True)
        assert workbook.active["A4"].value == 30
        workbook.close()


def test_arithmetic_shows_steps_and_rejects_code():
    result, steps = calculate("(12 + 18) * 2 / 3")
    assert result == 20
    assert len(steps) == 3
    for expression in ("__import__('os').system('id')", "2 ** 10000", "1 / 0", "1e309"):
        with pytest.raises((ValueError, ZeroDivisionError)):
            calculate(expression)


def test_spreadsheet_cycles_and_external_formulas_are_rejected(tmp_path):
    for rows in ([["=A1"]], [["=WEBSERVICE('https://example.com')"]]):
        with pytest.raises(ValueError):
            write_xlsx(tmp_path / "bad.xlsx", rows=rows)
    assert not (tmp_path / "bad.xlsx").exists()


def test_document_search_citations_and_deletion(client):
    uploaded = client.post(
        "/api/documents",
        files={
            "file": (
                "manual.txt",
                b"Valve corrosion inspection requires replacement.",
                "text/plain",
            )
        },
    ).json()
    results = client.post("/api/documents/search", json={"query": "valve corrosion"}).json()
    assert results[0]["citation"]["document_id"] == uploaded["document_id"]
    assert results[0]["citation"]["span"] == "0:48"
    assert client.delete(f"/api/documents/{uploaded['document_id']}").status_code == 200
    assert client.post("/api/documents/search", json={"query": "valve corrosion"}).json() == []
    assert client.post("/api/tasks", json={"prompt": "   "}).status_code == 422


def test_docx_tables_are_searchable(tmp_path):
    path = tmp_path / "table.docx"
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Pump vibration exceeds tolerance"
    document.save(path)
    assert "Pump vibration" in extract_pages(path)[0][1]


@pytest.mark.skipif(
    not shutil.which("tesseract") or not shutil.which("pdftoppm"), reason="OCR binaries required"
)
def test_real_scanned_pdf_to_approval_note(client, monkeypatch, tmp_path):
    from sovereign_workbench.core.llm.registry import get_registry

    monkeypatch.setattr(get_registry(), "probed", True)
    image = Image.new("RGB", (1400, 400), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("DejaVuSans.ttf", 32)
    draw.text((35, 50), "Inspection report R17", font=font, fill="black")
    draw.text((35, 120), "Valve corrosion found. Replace before restart.", font=font, fill="black")
    path = tmp_path / "scan.pdf"
    image.save(path, "PDF", resolution=150)
    image.close()
    upload = client.post(
        "/api/documents", files={"file": ("scan.pdf", path.read_bytes(), "application/pdf")}
    )
    assert upload.json()["chunks"] > 0
    response = client.post(
        "/api/chat",
        json={
            "prompt": "Draft an approval note from the scanned inspection report",
            "attachments": [upload.json()["document_id"]],
        },
    )
    assert "event: done" in response.text, response.text
    artifact = client.get("/api/artifacts").json()[0]
    content = client.get(f"/api/artifacts/{artifact['id']}").content
    text = "\n".join(p.text for p in Document(BytesIO(content)).paragraphs)
    assert "Valve corrosion" in text
    assert client.get("/api/audit/egress").json()["external_allowed"] == 0


def test_sandbox_timeout_and_output_bounds():
    async def run():
        sandbox = SubprocessSandbox()
        result = await sandbox.run("print('x' * 1000000)")
        assert result.exit_code == 0
        assert len(result.stdout) == 16_000
        result = await sandbox.run("import time; time.sleep(10)", timeout_s=1)
        assert result.timed_out
        assert result.exit_code != 0

    asyncio.run(run())
