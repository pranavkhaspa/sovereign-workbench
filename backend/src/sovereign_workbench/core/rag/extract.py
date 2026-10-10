"""Text extraction from local files. Heavy parsers are imported lazily."""

from __future__ import annotations

from pathlib import Path

from sovereign_workbench.observability.logging import get_logger

log = get_logger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".log", ".json", ".yaml", ".yml"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


def extract_pages(path: Path) -> list[tuple[int | None, str]]:
    """Return a list of (page_number, text). Page is None for non-paginated formats."""
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return [(None, path.read_text(encoding="utf-8", errors="replace"))]
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _extract_xlsx(path)
    if suffix in IMAGE_SUFFIXES:
        from sovereign_workbench.core.multimodal.ocr import extract_text

        return [(1, extract_text(path))]
    log.warning("unsupported file type", extra={"extra_fields": {"suffix": suffix}})
    return []


def _extract_pdf(path: Path) -> list[tuple[int | None, str]]:
    try:
        from pypdf import PdfReader  # lazy
    except ImportError:
        log.warning("pypdf not installed; falling back to OCR")
        from sovereign_workbench.core.multimodal.ocr import extract_text

        return [(1, extract_text(path))]
    reader = PdfReader(str(path))
    pages: list[tuple[int | None, str]] = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append((number, text))
        else:
            pages.append((number, ""))
    if any(text.strip() for _, text in pages):
        return pages
    from sovereign_workbench.core.multimodal.ocr import extract_pdf_pages

    ocr_pages = extract_pdf_pages(path)
    if ocr_pages:
        return [(index, text) for index, text in enumerate(ocr_pages, start=1)]
    return pages


def _extract_docx(path: Path) -> list[tuple[int | None, str]]:
    try:
        import docx  # lazy
    except ImportError:
        log.warning("python-docx not installed")
        return []
    document = docx.Document(str(path))
    return [(None, "\n".join(p.text for p in document.paragraphs))]


def _extract_xlsx(path: Path) -> list[tuple[int | None, str]]:
    try:
        from openpyxl import load_workbook  # lazy
    except ImportError:
        log.warning("openpyxl not installed")
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    pages: list[tuple[int | None, str]] = []
    for sheet in workbook.worksheets:
        lines = []
        for row in sheet.iter_rows(values_only=True):
            lines.append(" | ".join("" if c is None else str(c) for c in row))
        pages.append((None, f"# sheet: {sheet.title}\n" + "\n".join(lines)))
    return pages
