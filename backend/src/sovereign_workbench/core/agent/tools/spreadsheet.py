"""Spreadsheet reading tool (openpyxl)."""

from __future__ import annotations

from pathlib import Path

from pydantic import BaseModel, Field

from sovereign_workbench.core.agent.tools.base import Tool, ToolContext, ToolResult
from sovereign_workbench.core.agent.tools.fs_tools import _safe_path


class ReadSpreadsheetArgs(BaseModel):
    path: str = Field(description=".xlsx path relative to the workspace.")
    sheet: str | None = Field(default=None, description="Sheet name; default is the first sheet.")
    max_rows: int = Field(default=50, ge=1, le=1000)


class ReadSpreadsheetTool(Tool):
    name = "read_spreadsheet"
    description = "Read an .xlsx sheet and return its rows as text."
    args_schema = ReadSpreadsheetArgs

    async def run(self, args: BaseModel, ctx: ToolContext) -> ToolResult:
        assert isinstance(args, ReadSpreadsheetArgs)
        try:
            from openpyxl import load_workbook  # lazy
        except ImportError:
            return ToolResult(ok=False, error="openpyxl is not installed")
        try:
            target: Path = _safe_path(ctx, args.path)
            workbook = load_workbook(target, read_only=True, data_only=True)
            sheet = workbook[args.sheet] if args.sheet else workbook[workbook.sheetnames[0]]
            rows = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= args.max_rows:
                    break
                rows.append(" | ".join("" if c is None else str(c) for c in row))
            return ToolResult(
                ok=True,
                output=f"sheet={sheet.title} rows={len(rows)}\n" + "\n".join(rows),
                data={"sheet": sheet.title, "rows": len(rows)},
            )
        except Exception as exc:  # noqa: BLE001
            return ToolResult(ok=False, error=str(exc))
