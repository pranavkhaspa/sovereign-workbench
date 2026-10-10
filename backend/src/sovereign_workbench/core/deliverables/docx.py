"""Word document writer (approval notes, reports)."""

from __future__ import annotations

from pathlib import Path

from sovereign_workbench.observability.logging import get_logger

log = get_logger(__name__)


def write_docx(
    destination: Path,
    *,
    title: str,
    paragraphs: list[str],
    bullets: list[str] | None = None,
    citations: list[str] | None = None,
) -> Path:
    from docx import Document  # lazy

    document = Document()
    document.add_heading(title, level=0)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if bullets:
        for item in bullets:
            document.add_paragraph(item, style="List Bullet")
    if citations:
        document.add_heading("References", level=2)
        for cite in citations:
            document.add_paragraph(cite, style="List Number")
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    log.info("wrote docx", extra={"extra_fields": {"path": str(destination)}})
    return destination
