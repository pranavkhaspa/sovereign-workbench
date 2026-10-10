"""Excel writer (computed spreadsheets)."""

from __future__ import annotations

import re
from collections.abc import Sequence
from pathlib import Path

from sovereign_workbench.core.deliverables.calculations import calculate
from sovereign_workbench.observability.logging import get_logger

log = get_logger(__name__)


def write_xlsx(
    destination: Path,
    *,
    sheet_name: str = "Sheet1",
    rows: Sequence[Sequence[object]],
) -> Path:
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    from xlsxwriter import Workbook  # lazy

    computed: dict[tuple[int, int], float] = {}
    visiting: set[tuple[int, int]] = set()

    def resolve(reference: str) -> float:
        row, column = coordinate_to_tuple(reference.upper().replace("$", ""))
        if row < 1 or column < 1:
            raise ValueError("invalid cell reference")
        key = (row, column)
        if key in computed:
            return computed[key]
        if key in visiting:
            raise ValueError(f"circular formula at {reference}")
        visiting.add(key)
        value = (
            rows[row - 1][column - 1] if row <= len(rows) and column <= len(rows[row - 1]) else 0
        )
        if isinstance(value, str) and value.startswith("="):
            formula = value[1:].replace("$", "").replace("^", "**")

            def expand(match: re.Match[str]) -> str:
                start_row, start_col = coordinate_to_tuple(match[1])
                end_row, end_col = coordinate_to_tuple(match[2])
                if (end_row - start_row + 1) * (end_col - start_col + 1) > 10_000:
                    raise ValueError("formula range is too large")
                return ",".join(
                    f"{get_column_letter(c)}{r}"
                    for r in range(start_row, end_row + 1)
                    for c in range(start_col, end_col + 1)
                )

            formula = re.sub(r"([A-Z]+[1-9][0-9]*):([A-Z]+[1-9][0-9]*)", expand, formula.upper())
            result, _ = calculate(formula, resolve)
        else:
            if not isinstance(value, (str, int, float, bool, type(None))):
                raise ValueError("unsupported cell value")
            result = float(value or 0)
        visiting.remove(key)
        computed[key] = result
        return result

    for row_number, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            if isinstance(value, str) and value.startswith("="):
                resolve(f"{get_column_letter(column)}{row_number}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with Workbook(destination, {"strings_to_urls": False}) as workbook:
        sheet = workbook.add_worksheet(sheet_name)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if isinstance(value, str) and value.startswith("="):
                    sheet.write_formula(r, c, value, None, computed[(r + 1, c + 1)])
                else:
                    sheet.write(r, c, value)
    log.info("wrote xlsx", extra={"extra_fields": {"path": str(destination)}})
    return destination
