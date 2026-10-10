"""PowerPoint writer (board / review decks)."""

from __future__ import annotations

from pathlib import Path

from sovereign_workbench.observability.logging import get_logger

log = get_logger(__name__)


def write_pptx(destination: Path, *, title: str, slides: list[dict[str, str]]) -> Path:
    from pptx import Presentation  # lazy
    from pptx.util import Inches

    presentation = Presentation()
    layout = presentation.slide_layouts[0]
    first = presentation.slides.add_slide(layout)
    first.shapes.title.text = title
    if first.placeholders[1] is not None:
        first.placeholders[1].text = "Generated on-premise by Sovereign Workbench"

    content_layout = presentation.slide_layouts[1]
    for slide in slides:
        page = presentation.slides.add_slide(content_layout)
        page.shapes.title.text = slide.get("title", "")
        body = page.placeholders[1].text_frame
        body.text = slide.get("body", "")
        body.width = Inches(8)

    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(destination))
    log.info("wrote pptx", extra={"extra_fields": {"path": str(destination)}})
    return destination
